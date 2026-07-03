# -*- coding: utf-8 -*-
"""XLSX loader for Clipping Test cases.

The workbook uses Modbus document addresses. The runner converts them to
pymodbus addresses by subtracting 1.
"""

from __future__ import annotations

import os


TYPE_WORDS = {
    "uint16": 1,
    "int16": 1,
    "uint32": 2,
    "int32": 2,
    "uint64": 4,
    "int64": 4,
    "float": 2,
}


def _default_xlsx_path(product: str) -> str:
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(os.path.dirname(here))
    return os.path.join(root, "config", f"clipping_test_{product.lower()}.xlsx")


def _text(value, default: str = "") -> str:
    if value is None:
        return default
    return str(value).strip()


def _parse_bool(value, default: bool = False) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    if isinstance(value, str):
        v = value.strip().lower()
        if not v:
            return default
        return v in ("true", "1", "yes", "y", "ok", "run", "enabled")
    return default


def _is_enabled(value) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        v = value.strip().lower()
        if v in ("", "false", "0", "no", "n", "disabled", "skip"):
            return False
        return True
    return _parse_bool(value)


def _parse_int(value, default=None):
    if value is None:
        return default
    if isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str):
        v = value.strip()
        if not v:
            return default
        return int(float(v))
    return default


def _parse_number(value, default=None):
    if value is None:
        return default
    if isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        v = value.strip()
        if not v:
            return default
        try:
            if "." in v or "e" in v.lower():
                return float(v)
            return int(v)
        except ValueError:
            return default
    return default


def _doc_address(value):
    addr = _parse_int(value)
    if addr is None:
        return None
    if addr <= 0:
        raise ValueError(f"invalid document address: {value!r}")
    return addr


def _value_type(value, default: str = "uint16") -> str:
    v = _text(value, default).lower()
    return v if v in TYPE_WORDS else default


def _case_from_record(record: dict, row_number: int, selected_product: str) -> dict | None:
    name = _text(record.get("name"))
    if not name or not _is_enabled(record.get("enabled")):
        return None

    product = _text(record.get("product"), selected_product) or selected_product
    value_type = _value_type(record.get("value_type"))
    read_value_type = _value_type(record.get("read_value_type"), value_type)
    words = _parse_int(record.get("words"), TYPE_WORDS.get(value_type, 1))
    read_words = TYPE_WORDS.get(read_value_type, words)

    addr_doc = _doc_address(record.get("addr_doc"))
    if addr_doc is None:
        raise ValueError(f"row {row_number}: addr_doc is required")
    read_doc = _doc_address(record.get("read_doc_addr")) or addr_doc
    access_doc = _doc_address(record.get("access_doc_addr"))

    return {
        "name": name,
        "row": row_number,
        "product": product,
        "access_doc_addr": access_doc,
        "access_addr": access_doc - 1 if access_doc is not None else None,
        "addr_doc": addr_doc,
        "addr": addr_doc - 1,
        "words": words,
        "value_type": value_type,
        "write_value": _parse_number(record.get("write_value")),
        "expected_value": _parse_number(record.get("expected_value")),
        "valid_low": _parse_number(record.get("valid_low")),
        "valid_high": _parse_number(record.get("valid_high")),
        "read_doc_addr": read_doc,
        "read_addr": read_doc - 1,
        "read_words": read_words,
        "read_value_type": read_value_type,
        "commit_value": _parse_number(record.get("commit_value"), 1),
        "settle_ms": _parse_int(record.get("settle_ms"), 500),
        "restore_value": _parse_number(record.get("restore_value")),
        "restore_after": _parse_bool(record.get("restore_after"), False),
        "notes": record.get("notes"),
    }


def load_clipping_cases(product: str = "A2700", xlsx_path: str | None = None) -> list[dict]:
    from openpyxl import load_workbook

    path = xlsx_path or _default_xlsx_path(product)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"clipping test xlsx not found: {path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if len(rows) < 3:
        return []

    header = [_text(h) for h in rows[0]]
    cases: list[dict] = []
    for row_number, row in enumerate(rows[2:], start=3):
        if row is None or all(v is None for v in row):
            continue
        record = dict(zip(header, row))
        case = _case_from_record(record, row_number, product)
        if case is not None:
            cases.append(case)
    return cases
