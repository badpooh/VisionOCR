# -*- coding: utf-8 -*-
"""xlsx 기반 Setup 테스트 케이스 로더.

vision/config/setup_test_<product>.xlsx 를 읽어서 케이스 dict 리스트로 반환.

컬럼:
    name, enabled (TRUE/FALSE/init),
    addr           "주소,words"  예: "6001,1"
    access_addr    "주소,words" 또는 빈. 예: "6000,1"
    value_type     "uint16" 또는 "uint16; float"  (addr_type[; access_type])
    target_value   값 (정수/실수)
    main_menu_xy / side_menu_xy / data_view_xy  메뉴 터치 좌표
    password       TRUE/M 이면 기존 0000 Enter, TRUE; MHN 이면 MHN 좌표 시퀀스
    input_type     popup/number — 사람용 메모, 코드는 분기 안 함
    popup_xy       다중 좌표 — popup 옵션 또는 number 입력의 digit 좌표 시퀀스
    apply_xy       단일 좌표
    expected_text  콤마 구분, multiset strict
    roi_xy         "x1,y1,x2,y2" 단일. 비면 전체 화면
    verify_defaults TRUE 면 매 케이스 후 default 보존 검증
    notes

words 와 value_type 이 모순이어도 코드는 방어 안 함 — addr 의 words 로
read/write 길이 결정, value_type 으로 부호/float 변환만. 결과가 이상하면
사용자가 수정.
"""

from __future__ import annotations

import os

from demo_test.demo_xlsx_loader import (
    _parse_xy_list,
    _parse_bool,
    _parse_enabled,
    _parse_fixed_text,
)


def _parse_xy_single(s):
    """단일 좌표 'x,y' → (x, y). 빈 셀 None."""
    if s is None:
        return None
    if not isinstance(s, str):
        raise ValueError(f"xy 셀은 문자열이어야 함: {type(s).__name__}={s!r}")
    s = s.strip()
    if not s:
        return None
    parts = [p.strip() for p in s.split(",")]
    if len(parts) != 2:
        raise ValueError(f"xy 항목은 'x,y' 형식이어야 함: {s!r}")
    return (int(parts[0]), int(parts[1]))


def _parse_roi(s):
    """'x1,y1,x2,y2' → (x1, y1, x2, y2). 빈 셀 None."""
    if s is None:
        return None
    if not isinstance(s, str):
        return None
    s = s.strip()
    if not s:
        return None
    parts = [p.strip() for p in s.split(",")]
    if len(parts) != 4:
        raise ValueError(f"roi 는 'x1,y1,x2,y2' 형식이어야 함: {s!r}")
    return tuple(int(p) for p in parts)


def _parse_addr_tuple(s):
    """'6001,1' → (6001, 1). 빈 셀 None.

    words 생략 시 1 word 가정 — 'addr,1' 또는 그냥 '6001' 도 허용.
    """
    if s is None:
        return None
    if isinstance(s, (int, float)) and not isinstance(s, bool):
        return (int(s), 1)
    if not isinstance(s, str):
        return None
    s = s.strip()
    if not s:
        return None
    parts = [p.strip() for p in s.split(",")]
    if len(parts) == 1:
        return (int(parts[0]), 1)
    return (int(parts[0]), int(parts[1]))


def _parse_value_type(s):
    """'uint16' → ['uint16'], 'uint16; float' → ['uint16', 'float'].

    유효 값: uint16 / int16 / uint32 / int32 / uint64 / int64 / float
    (대소문자 무관).
    """
    if s is None:
        return []
    if not isinstance(s, str):
        return []
    s = s.strip()
    if not s:
        return []
    return [p.strip().lower() for p in s.split(";") if p.strip()]


def _parse_one_number(s):
    """단일 숫자 문자열 → int 또는 float."""
    s = s.strip()
    if not s:
        return None
    try:
        if "." in s or "e" in s.lower():
            return float(s)
        return int(s)
    except ValueError:
        return None


def _parse_target_value(s):
    """target_value 셀 → list. '0' → [0], '0; 1' → [0, 1].

    list[0] = addr 에 write 할 값
    list[1] = access 에 write 할 값 (없으면 default 1)
    """
    if s is None:
        return []
    if isinstance(s, bool):
        return []
    if isinstance(s, (int, float)):
        return [s]
    if not isinstance(s, str):
        return []
    s = s.strip()
    if not s:
        return []
    parts = [p.strip() for p in s.split(";") if p.strip()]
    out = []
    for p in parts:
        v = _parse_one_number(p)
        if v is not None:
            out.append(v)
    return out


def _parse_input_type(s):
    if s is None:
        return None
    if isinstance(s, str):
        v = s.strip().lower()
        if v in ("popup", "number"):
            return v
    return None


def _parse_password_mode(s):
    """Password mode parser.

    Backward compatible:
      TRUE / 1 / yes / ok -> "M"      (legacy 0000 Enter)
      TRUE; M / M         -> "M"
      TRUE; MHN / MHN     -> "MHN"    (Rootech0 coordinate sequence)
      FALSE / blank       -> None
    """
    if s is None:
        return None
    if isinstance(s, bool):
        return "M" if s else None
    if isinstance(s, (int, float)):
        return "M" if bool(s) else None
    if not isinstance(s, str):
        return None

    tokens = [p.strip().upper() for p in s.split(";") if p and p.strip()]
    if not tokens:
        return None
    if tokens[0].lower() in ("false", "0", "no", "n"):
        return None
    if "MHN" in tokens:
        return "MHN"
    if "M" in tokens:
        return "M"
    if tokens[0].lower() in ("true", "1", "yes", "y", "ok"):
        return "M"
    return None


def _parse_number_input(s):
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return str(s) if not isinstance(s, bool) else None
    if isinstance(s, str):
        s = s.strip()
        return s if s else None
    return None


def _xlsx_path(product: str) -> str:
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(os.path.dirname(here))
    return os.path.join(root, "config", f"setup_test_{product.lower()}.xlsx")


def _defaults_xlsx_path(product: str) -> str:
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(os.path.dirname(here))
    return os.path.join(root, "config", f"defaults_{product.lower()}.xlsx")


def load_defaults_cases(product: str = "A3700N") -> list:
    """defaults_<product>.xlsx 를 자동 로드 — config 폴더 고정.

    init row 가 호출됐을 때 사용. 파일 없거나 유효 케이스 없으면
    FileNotFoundError. 시트/컬럼 형식은 setup_test xlsx 와 동일하므로
    load_setup_cases 를 재사용.
    """
    path = _defaults_xlsx_path(product)
    if not os.path.isfile(path):
        raise FileNotFoundError(
            f"defaults xlsx 없음: {path}. "
            f"먼저 'Apply Defaults' 버튼으로 파일 만들거나 직접 작성하세요."
        )
    cases = load_setup_cases(product, xlsx_path=path)
    cases = [c for c in cases if c.get("addr") is not None]
    if not cases:
        raise FileNotFoundError(
            f"defaults xlsx 에 유효한 케이스 없음 (addr 비어있음): {path}"
        )
    return cases


def load_setup_cases(product: str = "A3700N", xlsx_path: str = None) -> list:
    from openpyxl import load_workbook

    path = xlsx_path if xlsx_path else _xlsx_path(product)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"setup test xlsx 없음: {path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if len(rows) < 3:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]

    cases = []
    for r in rows[2:]:
        if r is None or all(v is None for v in r):
            continue
        record = dict(zip(header, r))
        if not record.get("name"):
            continue
        state = _parse_enabled(record.get("enabled"))
        if state == "disabled":
            continue
        cases.append({
            "name":          str(record["name"]).strip(),
            "is_init":       (state == "init"),
            "addr":          _parse_addr_tuple(record.get("addr")),
            "access_addr":   _parse_addr_tuple(record.get("access_addr")),
            "value_type":    _parse_value_type(record.get("value_type")),
            "target_value":  _parse_target_value(record.get("target_value")),
            "main_menu_xy":  _parse_xy_list(record.get("main_menu_xy")),
            "side_menu_xy":  _parse_xy_list(record.get("side_menu_xy")),
            "data_view_xy":  _parse_xy_list(record.get("data_view_xy")),
            "password":      _parse_password_mode(record.get("password")),
            "input_type":    _parse_input_type(record.get("input_type")),
            "popup_xy":      _parse_xy_list(record.get("popup_xy")),
            "apply_xy":      _parse_xy_list(record.get("apply_xy")),
            "expected_text": _parse_fixed_text(record.get("expected_text")),
            "roi_xy":        _parse_roi(record.get("roi_xy")),
            "verify_defaults": _parse_bool(record.get("verify_defaults")),
            "notes":         record.get("notes"),
        })
    return cases
