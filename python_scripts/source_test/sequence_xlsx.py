from __future__ import annotations

import re


TEST_CASE_HEADERS = [
    "TC_ID", "Enable", "Test Name", "Feature", "Settle(s)", "Notes",
]
SETUP_HEADERS = [
    "TC_ID", "Order", "Name", "Address", "Access Address",
    "Value Type", "Write Value", "Access Value", "Notes",
]
CMC_HEADERS = [
    "TC_ID", "Order", "Name", "Duration(s)",
    "Va", "Vb", "Vc", "Ia", "Ib", "Ic", "Frequency", "Notes",
]
NAV_HEADERS = [
    "TC_ID", "Order", "Name", "button_keyin", "main_menu_xy",
    "side_menu_xy", "data_view_xy", "Wait(s)", "Notes",
]
EXPECTED_HEADERS = [
    "TC_ID", "Source Name", "Check Name", "Expected", "Unit",
    "Tolerance", "roi_xy", "Required Text", "Notes",
]

BUTTON_KEYINS = {
    "home": 0x01,
    "setup": 0x02,
    "event": 0x03,
    "back": 0x04,
    "esc": 0x04,
    "escape": 0x04,
}


def load_sequence(path: str) -> list[dict]:
    """Compatibility entry point used by the UI."""
    return load_test_plan(path)


def load_test_plan(path: str) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    cases = _read_test_cases(wb["TestCases"] if "TestCases" in wb.sheetnames else wb.active)
    by_id = {case["tc_id"]: case for case in cases}

    if "SetupModbus" in wb.sheetnames:
        for row in _read_rows(wb["SetupModbus"]):
            tc_id = _text(row, "TC_ID")
            if tc_id in by_id:
                by_id[tc_id]["setup_modbus"].append(_setup_from_row(row))

    if "CMC" in wb.sheetnames:
        for row in _read_rows(wb["CMC"]):
            tc_id = _text(row, "TC_ID")
            if tc_id in by_id:
                by_id[tc_id]["cmc_outputs"].append(_cmc_from_row(row))

    if "Navigation" in wb.sheetnames:
        for row in _read_rows(wb["Navigation"]):
            tc_id = _text(row, "TC_ID")
            if tc_id in by_id:
                by_id[tc_id]["navigation"].append(_navigation_from_row(row))

    if "Expected" in wb.sheetnames:
        for row in _read_rows(wb["Expected"]):
            tc_id = _text(row, "TC_ID")
            if tc_id in by_id:
                by_id[tc_id]["expected"].append(_expected_from_row(row))

    for case in cases:
        case["setup_modbus"].sort(key=lambda item: item.get("order", 0))
        case["cmc_outputs"].sort(key=lambda item: item.get("order", 0))
        case["navigation"].sort(key=lambda item: item.get("order", 0))
    return [case for case in cases if case.get("enabled", True)]


def save_sequence(path: str, cases: list[dict]):
    """Save a functional test plan workbook."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases"
    ws.append(TEST_CASE_HEADERS)
    setup = wb.create_sheet("SetupModbus")
    setup.append(SETUP_HEADERS)
    cmc = wb.create_sheet("CMC")
    cmc.append(CMC_HEADERS)
    nav = wb.create_sheet("Navigation")
    nav.append(NAV_HEADERS)
    exp = wb.create_sheet("Expected")
    exp.append(EXPECTED_HEADERS)

    for case in cases:
        tc_id = case.get("tc_id", "")
        ws.append([
            tc_id,
            case.get("enabled", True),
            case.get("name", ""),
            case.get("feature", ""),
            case.get("settle_s", 3),
            case.get("notes", ""),
        ])
        for item in case.get("setup_modbus", []):
            setup.append([
                tc_id,
                item.get("order", ""),
                item.get("name", ""),
                item.get("doc_address", ""),
                item.get("doc_access_address", ""),
                item.get("value_type", "uint16"),
                item.get("write_value", ""),
                item.get("access_value", 1),
                item.get("notes", ""),
            ])
        for item in case.get("cmc_outputs", []):
            cmc.append([
                tc_id,
                item.get("order", ""),
                item.get("name", ""),
                item.get("duration", ""),
                item.get("va", ""),
                item.get("vb", ""),
                item.get("vc", ""),
                item.get("ia", ""),
                item.get("ib", ""),
                item.get("ic", ""),
                item.get("frequency", ""),
                item.get("notes", ""),
            ])
        for item in case.get("navigation", []):
            nav.append([
                tc_id,
                item.get("order", ""),
                item.get("name", ""),
                _button_keyins_to_text(item.get("button_keyin")),
                _xy_to_text(item.get("main_menu_xy")),
                _xy_to_text(item.get("side_menu_xy")),
                _xy_to_text(item.get("data_view_xy")),
                item.get("wait_s", ""),
                item.get("notes", ""),
            ])
        for item in case.get("expected", []):
            exp.append([
                tc_id,
                item.get("source_name", ""),
                item.get("check_name", ""),
                item.get("expected", ""),
                item.get("unit", ""),
                item.get("tolerance", ""),
                _roi_to_text(item.get("roi_xy")),
                item.get("required_text", ""),
                item.get("notes", ""),
            ])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 16
    wb.save(path)


def _read_test_cases(ws) -> list[dict]:
    cases = []
    for row in _read_rows(ws):
        tc_id = _text(row, "TC_ID")
        if not tc_id:
            continue
        cases.append({
            "tc_id": tc_id,
            "enabled": _bool(row.get("Enable"), True),
            "name": _text(row, "Test Name", tc_id),
            "feature": _text(row, "Feature"),
            "settle_s": _float(row.get("Settle(s)"), 3),
            "notes": _text(row, "Notes"),
            "setup_modbus": [],
            "cmc_outputs": [],
            "navigation": [],
            "expected": [],
        })
    return cases


def _read_rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    out = []
    for values in rows[1:]:
        if not values or all(v is None or str(v).strip() == "" for v in values):
            continue
        row = {}
        for idx, key in enumerate(header):
            if key:
                row[key] = values[idx] if idx < len(values) else None
        out.append(row)
    return out


def _setup_from_row(row: dict) -> dict:
    doc_addr = _int(row.get("Address"))
    doc_access = _int(row.get("Access Address"))
    value_type = _text(row, "Value Type", "uint16").lower()
    return {
        "order": _float(row.get("Order"), 0),
        "name": _text(row, "Name", f"setup_{doc_addr}"),
        "doc_address": doc_addr,
        "address": doc_addr - 1 if doc_addr is not None else None,
        "doc_access_address": doc_access,
        "access_address": doc_access - 1 if doc_access is not None else None,
        "value_type": value_type,
        "write_value": _typed_value(row.get("Write Value"), value_type),
        "access_value": _typed_value(row.get("Access Value"), "uint16", default=1),
        "notes": _text(row, "Notes"),
    }


def _cmc_from_row(row: dict) -> dict:
    return {
        "order": _float(row.get("Order"), 0),
        "name": _text(row, "Name", "CMC Output"),
        "duration": _float(row.get("Duration(s)"), 0),
        "va": _float(row.get("Va"), 0),
        "vb": _float(row.get("Vb"), 0),
        "vc": _float(row.get("Vc"), 0),
        "ia": _float(row.get("Ia"), 0),
        "ib": _float(row.get("Ib"), 0),
        "ic": _float(row.get("Ic"), 0),
        "frequency": _float(row.get("Frequency"), 60),
        "notes": _text(row, "Notes"),
    }


def _navigation_from_row(row: dict) -> dict:
    return {
        "order": _float(row.get("Order"), 0),
        "name": _text(row, "Name", "Navigate"),
        "main_menu_xy": _parse_xy_list(row.get("main_menu_xy")),
        "side_menu_xy": _parse_xy_list(row.get("side_menu_xy")),
        "data_view_xy": _parse_xy_list(row.get("data_view_xy")),
        "button_keyin": _parse_button_keyins(row.get("button_keyin")),
        "wait_s": _float(row.get("Wait(s)"), 0),
        "notes": _text(row, "Notes"),
    }


def _expected_from_row(row: dict) -> dict:
    return {
        "source_name": _text(row, "Source Name"),
        "check_name": _text(row, "Check Name", "Check"),
        "expected": _float(row.get("Expected"), None),
        "unit": _text(row, "Unit"),
        "tolerance_type": "percent",
        "tolerance": _float(row.get("Tolerance"), 0),
        "roi_xy": _parse_roi(row.get("roi_xy")),
        "required_text": _text(row, "Required Text"),
        "notes": _text(row, "Notes"),
    }


def _text(row: dict, key: str, default="") -> str:
    value = row.get(key)
    if value is None:
        return default
    return str(value).strip()


def _bool(value, default=False) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("true", "1", "yes", "y", "on")


def _float(value, default=0.0):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _int(value):
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _typed_value(value, value_type: str, default=None):
    if value is None or value == "":
        return default
    if value_type == "float":
        return float(value)
    return int(float(value))


def _parse_xy_list(value) -> list[list[int]]:
    text = "" if value is None else str(value).strip()
    if not text:
        return []
    out = []
    for part in text.split(";"):
        nums = [n.strip() for n in part.split(",") if n.strip()]
        if len(nums) >= 2:
            out.append([int(float(nums[0])), int(float(nums[1]))])
    return out


def _parse_button_keyins(value) -> list[int]:
    text = "" if value is None else str(value).strip()
    if not text:
        return []
    out = []
    for part in re.split(r"[;,]", text):
        token = part.strip()
        if not token:
            continue
        key = re.sub(r"[\s_-]+", "", token.lower())
        if key in BUTTON_KEYINS:
            out.append(BUTTON_KEYINS[key])
            continue
        try:
            out.append(int(token, 0))
        except ValueError:
            out.append(int(float(token)))
    return out


def _parse_roi(value):
    text = "" if value is None else str(value).strip()
    if not text:
        return None
    nums = [int(float(n.strip())) for n in text.split(",") if n.strip()]
    return nums[:4] if len(nums) >= 4 else None


def _xy_to_text(value) -> str:
    return "; ".join(f"{xy[0]}, {xy[1]}" for xy in (value or []))


def _button_keyins_to_text(value) -> str:
    return "; ".join(f"0x{int(keyin) & 0xFF:02X}" for keyin in (value or []))


def _roi_to_text(value) -> str:
    return ", ".join(str(v) for v in value) if value else ""
